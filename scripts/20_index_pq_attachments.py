"""Stage 1: index every parliamentary-answer attachment we hold.

Attachments are heterogeneous (xlsx / docx / pdf) and each answers a different
question, so rather than force them into one schema this builds a browsable
index: which PQ, which topic, what the file's tables are actually about, and the
citable URL. That makes the collection usable without pre-judging which tables
matter.
"""
import csv, glob, json, pathlib, re

ROOT = pathlib.Path(__file__).resolve().parent.parent
RAW = ROOT / "data" / "raw" / "pq_attachments"
OUT = ROOT / "data" / "processed"

man = {}
for line in (ROOT / "data" / "raw" / "manifest.jsonl").read_text().splitlines():
    e = json.loads(line)
    if e.get("status") == "ok" and e["file"].startswith("pq_attachments/"):
        man[e["file"].split("/", 1)[1]] = e

pq_by_url = {}
for f in sorted(glob.glob(str(ROOT / "data/raw/oireachtas/pqs_*.json"))):
    for q in json.load(open(f)):
        pq_by_url[q["url"]] = q


def clean(t, n=220):
    return re.sub(r"\s+", " ", t or "").strip()[:n]


def probe_xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheets, head = [], []
    for sn in wb.sheetnames:
        ws = wb[sn]
        sheets.append(sn)
        for r in ws.iter_rows(max_row=3, values_only=True):
            head += [str(c) for c in r if c not in (None, "")]
    wb.close()
    return sheets, clean(" | ".join(head[:18]))


def probe_docx(path):
    import docx
    d = docx.Document(path)
    head = []
    for t in d.tables[:3]:
        for r in t.rows[:2]:
            head += [c.text.strip() for c in r.cells if c.text.strip()]
    txt = " ".join(p.text for p in d.paragraphs if p.text.strip())
    return [f"{len(d.tables)} table(s)"], clean(" | ".join(dict.fromkeys(head)) or txt)


def probe_pdf(path):
    import pdfplumber
    with pdfplumber.open(path) as pdf:
        pages = len(pdf.pages)
        txt = pdf.pages[0].extract_text() or ""
    return [f"{pages} page(s)"], clean(txt)


rows = []
for path in sorted(glob.glob(str(RAW / "*"))):
    name = pathlib.Path(path).name
    if name.startswith("_"):
        continue
    meta = man.get(name, {})
    pq = pq_by_url.get(meta.get("source_page", ""), {})
    ext = name.strip().rsplit(".", 1)[-1].lower()
    try:
        if ext == "xlsx":
            parts, summary = probe_xlsx(path)
        elif ext == "docx":
            parts, summary = probe_docx(path)
        elif ext == "pdf":
            parts, summary = probe_pdf(path)
        else:
            parts, summary = [], ""
    except Exception as exc:                                  # noqa: BLE001
        parts, summary = [], f"(could not read: {type(exc).__name__})"
    rows.append({
        "file": f"data/raw/pq_attachments/{name}",
        "format": ext,
        "bytes": meta.get("bytes", ""),
        "pq_date": pq.get("date", ""),
        "pq_number": pq.get("question_number", ""),
        "topic": pq.get("topic", ""),
        "asked_by": pq.get("asked_by", ""),
        "sheets_or_pages": "; ".join(parts),
        "content_summary": summary,
        "attachment_url": meta.get("url", ""),
        "pq_url": meta.get("source_page", ""),
    })

rows.sort(key=lambda r: (r["pq_date"], str(r["pq_number"])), reverse=True)
with open(OUT / "pq_attachments_index.csv", "w", newline="", encoding="utf-8") as fh:
    w = csv.DictWriter(fh, fieldnames=list(rows[0].keys()))
    w.writeheader()
    w.writerows(rows)
print(f"wrote pq_attachments_index.csv ({len(rows)} attachments)")

import collections
print("\nby topic:")
for t, n in collections.Counter(r["topic"] or "(unmatched)" for r in rows).most_common():
    print(f"  {n:4}  {t}")
print("\nby format:", dict(collections.Counter(r["format"] for r in rows)))
unread = [r for r in rows if r["content_summary"].startswith("(could not read")]
if unread:
    print(f"\n{len(unread)} unreadable: {[r['file'].rsplit('/',1)[-1] for r in unread]}")
