"""Stage 1 parsing: DETE employment permit XLSX -> tidy CSV.

The Department changed the spreadsheet layout several times over the decade:
  2015-2019  Year | Nationality | New | Renewal | Total | Refused | Withdrawn
             (the year's total sits on a row whose Nationality cell is blank
              or reads "Jan - Dec", NOT "Grand Total")
  2020-2023  year banner row, then Nationality[ (group)] | Issued | Refused | Withdrawn
  2024       Nationality | Issued | Refused | Withdrawn  (no banner row)
  2025-2026  year banner, then Nationality | Issued | Refused  (no Withdrawn column)
So we locate the header row by content, match columns by prefix, and pull each
file's own Grand Total row out separately to cross-check our summed figures.
"""
import csv, glob, os, re, pathlib
import openpyxl

ROOT = pathlib.Path(__file__).resolve().parent.parent
OUT = ROOT / "data" / "processed"
OUT.mkdir(parents=True, exist_ok=True)

TOTAL_LABELS = re.compile(r"^(grand\s*total|total|jan\s*-\s*dec)$", re.I)


def norm(v):
    return "" if v is None else str(v).replace("\xa0", " ").strip()


def num(v):
    s = norm(v).replace(",", "")
    if s == "":
        return None
    try:
        return int(float(s))
    except ValueError:
        return None


def find_header(ws):
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=12, values_only=True), start=1):
        if any(norm(c).lower().startswith("nationality") for c in row):
            return i, [norm(c) for c in row]
    return None, None


def col_map(hdr):
    """Match header cells by prefix so 'Nationality (group)' still resolves."""
    m = {}
    for j, h in enumerate(hdr):
        h = h.lower().strip()
        for key in ("nationality", "new", "renewal", "total", "issued", "refused", "withdrawn", "year"):
            if h.startswith(key) and key not in m:
                m[key] = j
    return m


rows_out, official = [], []
for path in sorted(glob.glob(str(ROOT / "data/raw/employment_permits/*/*nationality*.xlsx"))):
    year = int(re.search(r"(20\d\d)", os.path.basename(path)).group(1))
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    hdr_i, hdr = find_header(ws)
    if hdr_i is None:
        print(f"  !! no header row found in {path}")
        continue
    ci = col_map(hdr)
    has_withdrawn = "withdrawn" in ci
    rel = os.path.relpath(path, ROOT)

    for row in ws.iter_rows(min_row=hdr_i + 1, values_only=True):
        def g(key):
            j = ci.get(key)
            return num(row[j]) if j is not None and j < len(row) else None

        nat = norm(row[ci["nationality"]]) if ci.get("nationality") is not None else ""
        year_cell = norm(row[ci["year"]]) if ci.get("year") is not None else ""

        issued = g("issued")
        if issued is None:
            issued = g("total")
        new, renew = g("new"), g("renewal")
        if issued is None and (new is not None or renew is not None):
            issued = (new or 0) + (renew or 0)
        if issued is None and g("refused") is None:
            continue                                   # blank/decorative row

        rec = {"year": year, "nationality": nat, "issued": issued, "new": new,
               "renewal": renew, "refused": g("refused"),
               "withdrawn": g("withdrawn") if has_withdrawn else None,
               "source_file": rel}

        # The year's own total row: labelled "Grand Total"/"Jan - Dec", or (in the
        # 2015-2019 layout) identified by a 4-digit year in the Year column.
        # Rows with a blank nationality and a month in the Year column (e.g. "Dec"
        # in the 2017 file) are monthly sub-totals and are dropped.
        if TOTAL_LABELS.match(nat) or re.fullmatch(r"20\d\d", year_cell):
            rec.pop("nationality")
            rec["withdrawn_reported"] = has_withdrawn
            official.append(rec)
        elif nat:
            rows_out.append(rec)
    wb.close()

with open(OUT / "employment_permits_by_nationality.csv", "w", newline="") as fh:
    w = csv.DictWriter(fh, fieldnames=["year", "nationality", "issued", "new", "renewal",
                                       "refused", "withdrawn", "source_file"])
    w.writeheader()
    w.writerows(rows_out)
print(f"wrote employment_permits_by_nationality.csv  ({len(rows_out)} rows)")

with open(OUT / "employment_permits_annual_totals.csv", "w", newline="") as fh:
    w = csv.DictWriter(fh, fieldnames=["year", "issued", "new", "renewal", "refused",
                                       "withdrawn", "withdrawn_reported", "source_file"])
    w.writeheader()
    w.writerows(sorted(official, key=lambda r: r["year"]))
print(f"wrote employment_permits_annual_totals.csv  ({len(official)} rows)")

# Cross-check: our sum of per-nationality rows vs the file's own Grand Total.
print("\nyear  official  summed   delta   refused(official)")
summed = {}
for r in rows_out:
    s = summed.setdefault(r["year"], {"issued": 0, "refused": 0})
    for k in s:
        if r[k]:
            s[k] += r[k]
for o in sorted(official, key=lambda r: r["year"]):
    y = o["year"]
    s = summed.get(y, {"issued": 0, "refused": 0})
    d = (o["issued"] or 0) - s["issued"]
    flag = "" if d == 0 else "  <-- MISMATCH"
    print(f"{y}  {o['issued'] or 0:8,} {s['issued']:8,} {d:7,}   {o['refused'] or 0:6,}{flag}")
