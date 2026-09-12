"""Stage 1 parsing: employment permits BY PERMIT TYPE.

Parses the spreadsheets attached to parliamentary answers (see script 15), which
carry the permit-type dimension absent from DETE's annual publication:
  sheet "Nationality"     : Permit Type | Nationality | <year> Issued ...
  sheet "Economic Sector" : Economic Sector | Permit Type | <year> Issued ...
The first column is only written on the first row of each group, so it is
forward-filled. Each output row keeps the PQ it came from.
"""
import csv, json, pathlib, re
import openpyxl

ROOT = pathlib.Path(__file__).resolve().parent.parent
RAW = ROOT / "data" / "raw" / "pq_attachments"
OUT = ROOT / "data" / "processed"

man = {}
mf = ROOT / "data" / "raw" / "manifest.jsonl"
for line in mf.read_text().splitlines():
    e = json.loads(line)
    if e.get("status") == "ok" and e["file"].startswith("pq_attachments/"):
        man[e["file"].split("/", 1)[1]] = e


def norm(v):
    return "" if v is None else str(v).replace("\xa0", " ").strip()


def num(v):
    s = norm(v).replace(",", "")
    if s == "":
        return None
    try:
        return int(float(s))
    except ValueError:
        return None


nat_rows, sec_rows = [], []
for path in sorted(RAW.glob("*.xlsx")):
    meta = man.get(path.name, {})
    src = {"source_file": f"data/raw/pq_attachments/{path.name}",
           "source_url": meta.get("url", ""), "pq_url": meta.get("source_page", ""),
           "pq_note": meta.get("note", "")}
    wb = openpyxl.load_workbook(path, data_only=True)
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        rows = [[norm(c) for c in r] for r in ws.iter_rows(values_only=True)]
        if len(rows) < 3:
            continue
        years = [(j, y) for j, y in enumerate(rows[0]) if re.fullmatch(r"20\d\d", y)]
        hdr = [h.lower() for h in rows[1]]
        if not years or len(hdr) < 2:
            continue
        a_lbl, b_lbl = hdr[0], hdr[1]
        cur_a = ""
        for r in rows[2:]:
            if not any(r):
                continue
            a = r[0] or cur_a            # forward-fill the grouping column
            cur_a = a
            b = r[1] if len(r) > 1 else ""
            if a.lower().startswith("grand total") or not b:
                continue
            for j, y in years:
                v = num(r[j]) if j < len(r) else None
                if v is None:
                    continue
                rec = {"year": int(y), "value": v, **src}
                if "permit type" in a_lbl:
                    nat_rows.append({"permit_type": a, "nationality": b, **rec})
                elif "sector" in a_lbl:
                    sec_rows.append({"economic_sector": a, "permit_type": b, **rec})
    wb.close()

for name, rows, key in (("employment_permits_by_type_nationality.csv", nat_rows, "nationality"),
                        ("employment_permits_by_type_sector.csv", sec_rows, "economic_sector")):
    if not rows:
        print(f"(no rows for {name})")
        continue
    cols = [c for c in rows[0]]
    with open(OUT / name, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=cols)
        w.writeheader()
        w.writerows(rows)
    print(f"wrote {name}  ({len(rows):,} rows)")

# Totals by permit type and year. Aggregated PER SOURCE FILE first: several
# attachments can cover the same year, and summing across them would double
# count. Where years overlap we keep the most recent PQ's version.
import collections
per_src = collections.defaultdict(int)
for r in nat_rows:
    per_src[(r["source_file"], r["year"], r["permit_type"])] += r["value"]

pq_date = {}
for r in nat_rows:
    m = re.search(r"(\d{4}-\d{2}-\d{2})", r["source_file"] + " " + r.get("pq_note", ""))
    pq_date[r["source_file"]] = m.group(1) if m else ""

# for each year pick the source with the latest PQ date
best_src = {}
for (src, y, _t) in per_src:
    if y not in best_src or pq_date.get(src, "") > pq_date.get(best_src[y], ""):
        best_src[y] = src

rows_out = []
for (src, y, t), v in sorted(per_src.items()):
    rows_out.append({"year": y, "permit_type": t, "issued": v, "source_file": src,
                     "pq_date": pq_date.get(src, ""),
                     "preferred_for_year": src == best_src.get(y)})
if rows_out:
    with open(OUT / "employment_permits_by_type_annual.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=["year", "permit_type", "issued", "source_file",
                                           "pq_date", "preferred_for_year"])
        w.writeheader()
        w.writerows(rows_out)
    print(f"wrote employment_permits_by_type_annual.csv ({len(rows_out)} rows, "
          f"{len({r['source_file'] for r in rows_out})} source documents)")

    dete = {int(r["year"]): int(r["issued"])
            for r in csv.DictReader(open(OUT / "employment_permits_annual_totals.csv"))
            if r["issued"]}
    print(f"\nreconciliation of the preferred source for each year vs DETE annual totals:")
    print(f"{'year':6}{'by-type sum':>13}{'DETE total':>12}{'diff':>8}  source")
    for y in sorted(best_src):
        tot = sum(v for (src, yy, _), v in per_src.items() if yy == y and src == best_src[y])
        d = tot - dete.get(y, 0)
        print(f"{y:<6}{tot:13,}{dete.get(y, 0):12,}{d:+8,}  {pathlib.Path(best_src[y]).name}")
